"""Build the latest proposal: the Section 6 file plus Section 7 (design methodology, prototyping, validation,
hardware block diagram) and Section 8 (literature review: articles and patents).

python tools/build_proposal.py
    reads  docs/ECE_414_Proposal_Draft_with_Section6.docx
    writes docs/ECE_414_Proposal_Draft_with_Section8.docx (+ .pdf, with LibreOffice)

The section text and tables are in tools/section7_content.py and tools/section8_content.py; each may also update
earlier sections (update_earlier_sections) so that they stay consistent. Every paragraph, caption and table is built from
copies of the proposal's own formatting (body text, Fig./Table captions, three-line tables), headings get TOC
bookmarks and TOC lines, and citations written as {cite:key} are numbered from 30 in order of first appearance and
appended to Section 12. The table of contents page numbers are then refreshed from a LibreOffice rendering.
Run docs/figures/make_section7_figures.py first.
"""
import copy
import os
import re
import shutil
import subprocess
import sys
import tempfile

import docx
from docx.shared import Inches, Pt
from docx.oxml.ns import qn
from lxml import etree

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, "tools"))
import section7_content as C7  # noqa: E402
import section8_content as C8  # noqa: E402

SECTIONS = (C7, C8)
REFERENCES = {k: v for c in SECTIONS for k, v in c.REFERENCES.items()}
FIRST_FIG, FIRST_TAB = 5, 9  # Sections 3-6 have Figs. 1-4 and Tables 1-8

SRC = os.path.join(ROOT, "docs", "ECE_414_Proposal_Draft_with_Section6.docx")
DST = os.path.join(ROOT, "docs", "ECE_414_Proposal_Draft_with_Section8.docx")
W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
FIRST_NEW_REF = 30


def sub(parent, tag, **attrs):
    e = etree.SubElement(parent, qn(tag))
    for k, v in attrs.items():
        e.set(qn(k), str(v))
    return e


class Builder:
    def __init__(self, doc):
        self.doc = doc
        P = doc.paragraphs
        find = lambda pred: next(p for p in P if pred(p))
        self.t_h1 = find(lambda p: p.style.name == "Intro Section" and p.text.startswith("6."))._p
        self.t_h2 = find(lambda p: p.style.name == "Intro Subsection" and p.text.startswith("6.1"))._p
        self.t_body = find(lambda p: p.text.startswith("This section lists the engineering standards"))._p
        self.t_figcap = find(lambda p: p.text.startswith("Fig. 2."))._p
        self.t_figpara = find(lambda p: p._p.xpath(".//w:drawing") and P[P.index(p) + 1].text.startswith("Fig. 2."))._p
        self.t_tabcap = find(lambda p: p.text.startswith("Table 6."))._p
        self.t_ref = find(lambda p: p.text.startswith("9.\tMoryossef"))._p
        self.refs_heading = find(lambda p: p.style.name == "Intro Section" and p.text.startswith("12. References"))._p
        self.anchor = self.refs_heading.getprevious()  # the page-break paragraph before References
        assert self.anchor.xpath(".//w:br[@w:type='page']"), "expected a page break before References"
        self.toc_refs = find(lambda p: p.style.name == "toc 1" and p.text.startswith("12. References"))._p
        self.t_toc1 = self.toc_refs
        self.t_toc2 = find(lambda p: p.style.name == "toc 2")._p
        ids = [int(b.get(qn("w:id"))) for b in doc.element.body.iter(qn("w:bookmarkStart"))]
        self.next_bm = max(ids) + 1
        self.cites = {}
        self.toc_entries = []
        self.n_fig = 0
        self.n_tab = 0

    # ---- text with **bold**, *italic*, {cite:key} and {sup:x}
    def cite_text(self, text):
        def num(m):
            keys = [k.strip() for k in m.group(1).split(",")]
            out = []
            for k in keys:
                if k not in REFERENCES:
                    raise KeyError(f"unknown reference {k}")
                if k not in self.cites:
                    self.cites[k] = FIRST_NEW_REF + len(self.cites)
                out.append(f"[{self.cites[k]}]")
            return ", ".join(out)
        text = re.sub(r"\{cite:([^}]+)\}", num, text)
        return re.sub(r"\{ref:(\d+)\}", lambda m: f"[{m.group(1)}]", text)

    def runs(self, p, text, size=None, base_bold=False, base_italic=False):
        text = self.cite_text(text)
        for part in re.split(r"(\*\*.+?\*\*|(?<!\*)\*[^*]+?\*(?!\*)|\{sup:[^}]+\})", text):
            if not part:
                continue
            bold, italic, sup = base_bold, base_italic, False
            if part.startswith("**"):
                part, bold = part[2:-2], True
            elif part.startswith("{sup:"):
                part, sup = part[5:-1], True
            elif part.startswith("*") and part.endswith("*") and len(part) > 1:
                part, italic = part[1:-1], not base_italic
            r = sub(p, "w:r")
            rpr = sub(r, "w:rPr")
            if bold:
                sub(rpr, "w:b"); sub(rpr, "w:bCs")
            if italic:
                sub(rpr, "w:i"); sub(rpr, "w:iCs")
            if size:
                sub(rpr, "w:sz", **{"w:val": size}); sub(rpr, "w:szCs", **{"w:val": size})
            if sup:
                sub(rpr, "w:vertAlign", **{"w:val": "superscript"})
            if len(rpr) == 0:
                r.remove(rpr)
            t = sub(r, "w:t"); t.text = part
            t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")

    def para_like(self, template):
        p = etree.Element(qn("w:p"))
        ppr = template.find(qn("w:pPr"))
        if ppr is not None:
            p.append(copy.deepcopy(ppr))
        return p

    def insert(self, e):
        self.anchor.addprevious(e)
        return e

    # ---- blocks
    def page_break(self):
        p = etree.Element(qn("w:p")); r = sub(p, "w:r"); sub(r, "w:br", **{"w:type": "page"})
        self.insert(p)

    def heading(self, level, text):
        t = self.t_h1 if level == 1 else self.t_h2
        p = self.para_like(t)
        name = f"_Toc7{self.next_bm:05d}"
        bs = sub(p, "w:bookmarkStart", **{"w:id": self.next_bm, "w:name": name})
        r = sub(p, "w:r"); tt = sub(r, "w:t"); tt.text = text
        sub(p, "w:bookmarkEnd", **{"w:id": self.next_bm})
        self.next_bm += 1
        self.toc_entries.append((level, text, name))
        self.insert(p)

    def item_heading(self, text):
        """Bold, unindented heading inside a subsection (one per reviewed article or patent); kept with the next
        block so that it never ends a page on its own."""
        p = self.para_like(self.t_body)
        ppr = p.find(qn("w:pPr"))
        ind = ppr.find(qn("w:ind"))
        if ind is not None:
            ppr.remove(ind)
        ppr.insert(0, etree.Element(qn("w:keepNext")))
        sp = ppr.find(qn("w:spacing"))
        sp.set(qn("w:before"), "240")
        jc = ppr.find(qn("w:jc"))
        if jc is not None:
            jc.set(qn("w:val"), "left")
        self.runs(p, text, size=24, base_bold=True)
        self.insert(p)

    def body(self, text, indent=True):
        p = self.para_like(self.t_body)
        if not indent:
            ind = p.find(qn("w:pPr")).find(qn("w:ind"))
            if ind is not None:
                ind.getparent().remove(ind)
        self.runs(p, text, size=24)
        self.insert(p)

    def figure(self, path, width_in, caption):
        self.n_fig += 1
        num = FIRST_FIG + self.n_fig - 1
        para = self.doc.add_paragraph()
        para.add_run().add_picture(path, width=Inches(width_in))
        p = para._p
        p.getparent().remove(p)
        old = p.find(qn("w:pPr"))
        if old is not None:
            p.remove(old)
        p.insert(0, copy.deepcopy(self.t_figpara.find(qn("w:pPr"))))
        cap = self.cite_text(caption)
        plain = re.sub(r"\*+", "", cap)
        for dp in p.iter("{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}docPr"):
            dp.set("title", f"Fig. {num}."); dp.set("descr", plain); dp.set("name", f"Fig. {num}.")
        self.insert(p)
        c = self.para_like(self.t_figcap)
        c.find(qn("w:pPr")).insert(0, etree.Element(qn("w:keepLines")))  # never split a caption across pages
        r = sub(c, "w:r"); rpr = sub(r, "w:rPr"); sub(rpr, "w:b"); sub(rpr, "w:bCs")
        t = sub(r, "w:t"); t.text = f"Fig. {num}. "; t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        self.runs(c, cap)
        self.insert(c)
        return num

    def table(self, caption, widths, header, rows, shade_rows=(), group_rows=(), bold_first_col=False,
              center_cols=(), total_row=False):
        self.n_tab += 1
        num = FIRST_TAB + self.n_tab - 1
        c = self.para_like(self.t_tabcap)
        r = sub(c, "w:r"); rpr = sub(r, "w:rPr"); sub(rpr, "w:b"); sub(rpr, "w:bCs")
        t = sub(r, "w:t"); t.text = f"Table {num}. "; t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        self.runs(c, caption)
        self.insert(c)
        assert sum(widths) == 9360, (caption, sum(widths))
        tbl = etree.Element(qn("w:tbl"))
        tpr = sub(tbl, "w:tblPr")
        sub(tpr, "w:tblW", **{"w:w": 9360, "w:type": "dxa"})
        sub(tpr, "w:tblLayout", **{"w:type": "fixed"})
        sub(tpr, "w:tblLook", **{"w:val": "04A0", "w:firstRow": 1, "w:lastRow": 0, "w:firstColumn": 1,
                                  "w:lastColumn": 0, "w:noHBand": 0, "w:noVBand": 1})
        grid = sub(tbl, "w:tblGrid")
        for w in widths:
            sub(grid, "w:gridCol", **{"w:w": w})

        def border(tcpr, top, bottom):
            b = sub(tcpr, "w:tcBorders")
            for side, val in (("top", top), ("left", None), ("bottom", bottom), ("right", None)):
                if val:
                    sub(b, f"w:{side}", **{"w:val": "single", "w:sz": val, "w:space": 0, "w:color": "000000"})
                else:
                    sub(b, f"w:{side}", **{"w:val": "none", "w:sz": 0, "w:space": 0, "w:color": "FFFFFF"})

        def cell(tr, w, text, top, bottom, hdr=False, keep=True, shade=None, span=1, bold=False, italic=False,
                 center=False):
            tc = sub(tr, "w:tc"); tcpr = sub(tc, "w:tcPr")
            sub(tcpr, "w:tcW", **{"w:w": w, "w:type": "dxa"})
            if span > 1:
                sub(tcpr, "w:gridSpan", **{"w:val": span})
            border(tcpr, top, bottom)
            if shade:
                sub(tcpr, "w:shd", **{"w:val": "clear", "w:color": "auto", "w:fill": shade})
            mar = sub(tcpr, "w:tcMar")
            v = 80 if hdr else 50
            for side, val in (("top", v), ("left", 80), ("bottom", v), ("right", 80)):
                sub(mar, f"w:{side}", **{"w:w": val, "w:type": "dxa"})
            for k, line in enumerate(text.split("\n")):
                p = sub(tc, "w:p"); ppr = sub(p, "w:pPr")
                if keep:
                    sub(ppr, "w:keepNext")
                if center:
                    sub(ppr, "w:jc", **{"w:val": "center"})
                self.runs(p, line, base_bold=hdr or bold, base_italic=italic)

        tr = sub(tbl, "w:tr"); trpr = sub(tr, "w:trPr"); sub(trpr, "w:cantSplit"); sub(trpr, "w:tblHeader")
        for i, (w, h) in enumerate(zip(widths, header)):
            cell(tr, w, h, 12, 4, hdr=True, center=i in center_cols)
        for ri, row in enumerate(rows):
            last = ri == len(rows) - 1
            tr = sub(tbl, "w:tr"); trpr = sub(tr, "w:trPr"); sub(trpr, "w:cantSplit")
            top = 4 if (total_row and last) else None
            bottom = 12 if last else None
            if ri in group_rows:
                cell(tr, sum(widths), row[0], top, bottom, keep=True, span=len(widths), italic=True, shade="F2F2F2")
                continue
            shade = "E6E6E6" if ri in shade_rows else None
            for ci, (w, text) in enumerate(zip(widths, row)):
                cell(tr, w, text, top, bottom, keep=not last, shade=shade,
                     bold=(bold_first_col and ci == 0) or (total_row and last), center=ci in center_cols)
        self.insert(tbl)
        # small gap after the table, as in the rest of the proposal
        gap = etree.Element(qn("w:p")); ppr = sub(gap, "w:pPr"); sub(ppr, "w:spacing", **{"w:after": 120})
        self.insert(gap)
        return num

    # ---- edits to the earlier sections (keep their formatting, change only the text)
    def table_after_caption(self, prefix):
        cap = next(p for p in self.doc.paragraphs if p.text.startswith(prefix))._p
        tbl = cap.getnext()
        assert tbl.tag == qn("w:tbl"), prefix
        return tbl

    def set_cell(self, tbl, row, col, text, expect=None):
        """Replace the text of one cell; expect = the start of the current text, as a safety check."""
        tc = tbl.findall(qn("w:tr"))[row].findall(qn("w:tc"))[col]
        old = "".join(t.text or "" for t in tc.iter(qn("w:t")))
        if expect is not None and not old.startswith(expect):
            raise ValueError(f"row {row} col {col} is {old[:60]!r}, expected {expect!r}")
        paras = tc.findall(qn("w:p"))
        p = self.para_like(paras[0])
        rpr = paras[0].find(".//" + qn("w:rPr"))
        for q in paras:
            tc.remove(q)
        tc.append(p)
        self.runs(p, text)
        if rpr is not None:  # keep the cell's run formatting (e.g. bold totals)
            for r in p.findall(qn("w:r")):
                if r.find(qn("w:rPr")) is None:
                    r.insert(0, copy.deepcopy(rpr))

    def insert_row(self, tbl, before, texts):
        """Insert a body row before row index `before`, copying the formatting of the row above it."""
        rows = tbl.findall(qn("w:tr"))
        new = copy.deepcopy(rows[before - 1])
        rows[before].addprevious(new)
        for col, text in enumerate(texts):
            self.set_cell(tbl, before, col, text)

    def replace_text(self, prefix, old, new):
        """Replace text in the paragraph that starts with `prefix`; the match may span several runs. The new text
        takes the formatting of the run where the match starts."""
        p = next(p for p in self.doc.paragraphs if p.text.startswith(prefix))._p
        ts = [t for t in p.iter(qn("w:t"))]
        full = "".join(t.text or "" for t in ts)
        start = full.find(old)
        if start < 0:
            raise ValueError(f"{old!r} not found in paragraph {prefix!r}")
        end, pos, placed = start + len(old), 0, False
        for t in ts:
            s0, s1 = pos, pos + len(t.text or "")
            pos = s1
            if s1 <= start or s0 >= end:
                continue
            a, b_ = max(start, s0) - s0, min(end, s1) - s0
            t.text = t.text[:a] + ("" if placed else new) + t.text[b_:]
            t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
            placed = True

    # ---- TOC and references
    def toc(self):
        for level, text, name in self.toc_entries:
            t = copy.deepcopy(self.t_toc1 if level == 1 else self.t_toc2)
            h = t.find(qn("w:hyperlink"))
            h.set(qn("w:anchor"), name)
            ts = h.findall(".//" + qn("w:t"))
            ts[0].text = text
            for extra in ts[1:-1]:
                extra.text = ""
            ts[-1].text = "0"
            instr = h.find(".//" + qn("w:instrText"))
            instr.text = f" PAGEREF {name} \\h "
            self.toc_refs.addprevious(t)

    def keep_intro_on_page_3(self):
        """The longer TOC fills page 2, so the stand-alone page break after it would leave page 3 blank and push
        the Introduction to page 4 (it must start on page 3). Drop the spacer and the break paragraph after the
        TOC and start the Introduction with 'page break before' instead."""
        intro = next(p for p in self.doc.paragraphs if p.style.name == "Intro Section" and p.text.startswith("3."))._p
        e = intro.getprevious()
        while e is not None and e.tag == qn("w:p") and not e.xpath("string(.)").strip() \
                and not e.xpath(".//w:fldChar") and not e.xpath(".//w:drawing"):
            prev = e.getprevious()
            e.getparent().remove(e)
            e = prev
        # with Sections 7 and 8 the TOC has 35 lines; the styles' 5 pt after each line would push it onto page 3
        for name in ("toc 1", "toc 2"):
            self.doc.styles[name].paragraph_format.space_after = Pt(1)
        ppr = intro.find(qn("w:pPr"))
        if ppr.find(qn("w:pageBreakBefore")) is None:
            # schema order: pStyle, keepNext, keepLines, pageBreakBefore, ...
            before = [c for c in ppr if c.tag in (qn("w:pStyle"), qn("w:keepNext"), qn("w:keepLines"))]
            ppr.insert(ppr.index(before[-1]) + 1 if before else 0, etree.Element(qn("w:pageBreakBefore")))

    def references(self):
        last = self.refs_heading
        for e in self.refs_heading.itersiblings():
            if e.tag == qn("w:p") and e.xpath("string(.)").strip()[:3].rstrip(".\t").isdigit():
                last = e
        for key, n in sorted(self.cites.items(), key=lambda kv: kv[1]):
            p = self.para_like(self.t_ref)
            r = sub(p, "w:r"); rpr = sub(r, "w:rPr")
            sub(rpr, "w:sz", **{"w:val": 24}); sub(rpr, "w:szCs", **{"w:val": 24})
            t = sub(r, "w:t"); t.text = f"{n}."
            sub(r, "w:tab")
            self.runs(p, REFERENCES[key], size=24)
            last.addnext(p)
            last = p


def heading_pages(pdf):
    """Printed page number of every text line in the rendered PDF (pages without a footer number are skipped)."""
    out = subprocess.run(["pdftotext", "-layout", pdf, "-"], capture_output=True, text=True, check=True).stdout
    pages = out.split("\f")
    found = {}
    for i, page in enumerate(pages, 1):
        nums = [l.strip() for l in page.splitlines() if l.strip()]
        printed = int(nums[-1]) if nums and nums[-1].isdigit() else None  # page number in the footer
        if printed is None:
            continue
        for line in page.splitlines():
            s = " ".join(line.split())
            if s and "...." not in s and not re.search(r"\s\d+$", s):
                found.setdefault(s, []).append(printed)
    return found


def render_pdf(docx_path, out_dir):
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", out_dir, docx_path],
                   check=True, capture_output=True, timeout=300)
    return os.path.join(out_dir, os.path.splitext(os.path.basename(docx_path))[0] + ".pdf")


def refresh_toc(docx_path, tmp):
    """Set every TOC line's page number from a LibreOffice rendering (Word users can also update the field)."""
    pdf = render_pdf(docx_path, tmp)
    found = heading_pages(pdf)
    d = docx.Document(docx_path)
    toc_paras = [p for p in d.paragraphs if p.style.name in ("toc 1", "toc 2")]
    missing = []
    for p in toc_paras:
        title = " ".join(p.text.split("\t")[0].split())
        if not title:
            continue
        cands = [pg for key, pgs in found.items() if key.startswith(title) or title.startswith(key) and len(key) > 25
                 for pg in pgs]
        cands = [c for c in cands if c > 2]  # skip the TOC itself
        if not cands:
            missing.append(title); continue
        ts = p._p.findall(".//" + qn("w:t"))
        if ts and ts[-1].text.strip().isdigit():
            ts[-1].text = str(min(cands))
    d.save(docx_path)
    return missing


def main():
    doc = docx.Document(SRC)
    b = Builder(doc)
    for c in SECTIONS:
        c.update_earlier_sections(b)
    for c in SECTIONS:
        c.build(b)
    b.toc()
    b.keep_intro_on_page_3()
    b.references()
    doc.save(DST)
    tmp = tempfile.mkdtemp()
    missing = refresh_toc(DST, tmp)
    if missing:
        print("TOC lines without a page:", missing)
    pdf = render_pdf(DST, tmp)
    shutil.copy(pdf, DST.replace(".docx", ".pdf"))
    print("wrote", DST, "and", DST.replace(".docx", ".pdf"))
    print("new references:", {k: v for k, v in sorted(b.cites.items(), key=lambda kv: kv[1])})


if __name__ == "__main__":
    main()
